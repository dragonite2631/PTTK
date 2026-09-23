import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DIAGRAMS_DIR = os.path.join(BASE_DIR, "diagrams")
DOCX_OUT = os.path.join(BASE_DIR, "N12 Nhóm 01.docx")

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1E3A8A"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=160, right=160)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(30, 41, 59)

    # Set column widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # Spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Navy #1E3A8A
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(29, 78, 216) # Royal Blue #1D4ED8
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(51, 65, 85) # Slate #334155
    return h

def add_body_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11.5)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11)
    r_text.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_diagram_image(doc, filename, caption_text):
    path = os.path.join(DIAGRAMS_DIR, filename)
    if not os.path.exists(path):
        print(f"Warning: image {path} does not exist!")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(path, width=Inches(6.2))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(10)
    run_cap.font.italic = True
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGBColor(71, 85, 105)

def build_full_docx():
    doc = docx.Document()

    # Page setup - Margins: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm
    section = doc.sections[0]
    section.top_margin = Inches(0.79)    # 2.0 cm
    section.bottom_margin = Inches(0.79) # 2.0 cm
    section.left_margin = Inches(1.18)   # 3.0 cm
    section.right_margin = Inches(0.79)  # 2.0 cm

    # Document Header block
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    r = p_inst.add_run("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG\nKHOA CÔNG NGHỆ THÔNG TIN")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 41, 59)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(14)
    r_line = p_line.add_run("------------------***------------------")
    r_line.font.bold = True
    r_line.font.color.rgb = RGBColor(148, 163, 184)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    r_t = p_title.add_run("BÀI KIỂM TRA ĐỊNH KỲ\nPHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG THÔNG TIN")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(16)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(190, 18, 60) # Dark Red / Crimson

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("ĐỀ TÀI: XÂY DỰNG HỆ THỐNG ĐẶT VÉ THAM GIA SỰ KIỆN TRỰC TUYẾN")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(30, 58, 138)

    # Meta table
    meta_headers = ["Thông tin", "Chi tiết"]
    meta_data = [
        ["Lớp học phần", "N12"],
        ["Mã số nhóm", "Nhóm 01 (Sinh viên có thể thay số nhóm tương ứng)"],
        ["Giảng viên giảng dạy", "TS. Nguyễn Đức Hiển"],
        ["Email tiếp nhận bài", "ndhien@hotmail.com"],
        ["Tên file nộp bài", "N12 Nhóm 01.docx"]
    ]
    add_styled_table(doc, meta_headers, meta_data, col_widths=[2.0, 4.2])

    # =========================================================================
    # PHẦN 1
    # =========================================================================
    add_heading_1(doc, "1. MÃ SỐ NHÓM, TÊN ĐỀ TÀI, HỌ TÊN CÁC THÀNH VIÊN – CHỨC NĂNG")
    
    add_body_p(doc, "Xây dựng Hệ thống đặt vé tham gia sự kiện trực tuyến (Online Event Ticket Booking System)", "Tên đề tài: ")
    add_body_p(doc, "Nhóm 01 - Lớp tín chỉ: N12", "Mã số nhóm: ")
    add_body_p(doc, "Hệ thống phục vụ toàn diện nhu cầu bán vé các sự kiện âm nhạc (Concert), thể thao, hội thảo và biểu diễn nghệ thuật. Hệ thống giải quyết bài toán tải lớn khi mở bán thông qua phòng chờ ảo (Virtual Waiting Queue), sơ đồ ghế trực quan theo thời gian thực (Interactive Real-time Seat Map), khóa giữ ghế tạm thời, thanh toán đa kênh và cấp phát vé điện tử QR Code chống gian lận.", "Mô tả tổng quan: ")

    add_body_p(doc, "Dưới đây là bảng phân công nhiệm vụ và trách nhiệm cụ thể của từng thành viên tương ứng với 5 phân hệ (module) chính của bài tập môn học:")

    members_headers = ["STT", "Họ và tên", "Mã sinh viên", "Vai trò", "Chức năng / Module phụ trách"]
    members_data = [
        ["1", "Nguyễn Văn A", "B21DCCN001", "Nhóm trưởng", "Module 1: Quản lý tài khoản người dùng\n- Đăng ký, đăng nhập tài khoản, bảo mật phân quyền\n- Đổi mật khẩu, khôi phục mật khẩu quên qua Email OTP\n- Chỉnh sửa hồ sơ cá nhân, cài đặt thông báo"],
        ["2", "Trần Thị B", "B21DCCN002", "Thành viên", "Module 2: Tìm kiếm & Xem sự kiện\n- Trang chủ hiển thị sự kiện nổi bật, thịnh hành\n- Tìm kiếm, lọc sự kiện đa tiêu chí (ngày, địa điểm, thể loại)\n- Chi tiết sự kiện, dàn nghệ sĩ biểu diễn, lưu yêu thích"],
        ["3", "Lê Văn C", "B21DCCN003", "Thành viên", "Module 3: Xếp hàng và chọn chỗ\n- Hàng đợi ảo thông minh (Virtual Queue) khi lưu lượng cao\n- Hiển thị sơ đồ ghế sân khấu thời gian thực (Real-time Seat Map)\n- Chọn vị trí ghế, khóa ghế tạm thời (Hold Timer 10-15 phút)"],
        ["4", "Phạm Thị D", "B21DCCN004", "Thành viên", "Module 4: Tính tiền và Xuất vé\n- Tạo đơn hàng mua vé, áp dụng mã voucher khuyến mãi\n- Tích hợp cổng thanh toán trực tuyến (VNPAY, Ví MoMo, Thẻ ATM/Visa)\n- Tạo vé điện tử mã QR Code động, gửi vé qua Email"],
        ["5", "Hoàng Văn E", "B21DCCN005", "Thành viên", "Module 5: Quản lý và thống kê dành cho ban tổ chức\n- Công cụ vẽ & thiết kế sơ đồ sân khấu/ghế ngồi\n- Tạo và đăng bài sự kiện, quản lý suất diễn, cấu hình giá vé\n- Quản lý check-in khách hàng, báo cáo thống kê doanh thu"]
    ]
    add_styled_table(doc, members_headers, members_data, col_widths=[0.5, 1.3, 1.1, 1.0, 2.6])

    # =========================================================================
    # PHẦN 2
    # =========================================================================
    add_heading_1(doc, "2. BIỂU ĐỒ USE CASE CHO HỆ THỐNG VÀ CHO CÁC CHỨC NĂNG")
    
    add_heading_2(doc, "2.1. Biểu đồ Use Case tổng thể toàn hệ thống")
    add_body_p(doc, "Biểu đồ Use Case tổng quát mô tả toàn bộ ranh giới của Hệ thống đặt vé tham gia sự kiện trực tuyến, bao gồm các tác nhân (Actors) chính và các nhóm chức năng tương ứng:")
    add_bullet_p(doc, "Người dùng tìm kiếm, chọn ghế, thanh toán vé và quản lý vé cá nhân.", "Khách hàng (Customer): ")
    add_bullet_p(doc, "Đơn vị tổ chức thiết kế sơ đồ ghế, đăng sự kiện, cấu hình giá vé và theo dõi doanh thu.", "Ban tổ chức (Organizer): ")
    add_bullet_p(doc, "Kiểm duyệt sự kiện, quản lý tài khoản người dùng và giám sát hệ thống.", "Quản trị viên (Administrator): ")
    add_bullet_p(doc, "Hệ thống thanh toán trực tuyến xử lý giao dịch tiền tệ (VNPAY, MoMo, Visa/Mastercard).", "Cổng thanh toán (Payment Gateway - External): ")
    add_bullet_p(doc, "Dịch vụ bên ngoài gửi mã xác thực OTP, email hóa đơn và vé điện tử đính kèm QR Code.", "Dịch vụ Email / SMS (Notification Service - External): ")

    add_diagram_image(doc, "uc_tong_the.png", "Hình 2.1: Biểu đồ Use Case tổng thể Hệ thống Đặt vé tham gia sự kiện")

    # Module 1
    add_heading_2(doc, "2.2. Biểu đồ Use Case Phân hệ 1: Quản lý tài khoản người dùng")
    add_body_p(doc, "Phân hệ đảm nhiệm việc xác thực và định danh người dùng trong toàn hệ thống. Hỗ trợ đầy đủ các nghiệp vụ đăng ký, đăng nhập an toàn, cơ chế khôi phục mật khẩu qua mã xác nhận OTP/Token, cập nhật hồ sơ và tùy chỉnh nhận thông báo.")
    add_diagram_image(doc, "uc_module1.png", "Hình 2.2: Biểu đồ Use Case Phân hệ 1 - Quản lý tài khoản người dùng")

    uc1_headers = ["Use Case", "Tác nhân", "Mô tả nghiệp vụ", "Quan hệ Use Case"]
    uc1_data = [
        ["Đăng ký tài khoản", "Khách hàng, Organizer", "Người dùng điền thông tin (họ tên, email, mật khẩu, SĐT) để tạo tài khoản mới.", "<<include>> Gửi mã OTP xác nhận email"],
        ["Đăng nhập", "Người dùng", "Xác thực bằng Email/Username và Mật khẩu. Hỗ trợ ghi nhớ phiên đăng nhập qua JWT.", "<<extend>> Khóa tài khoản tạm thời khi nhập sai quá 5 lần"],
        ["Đổi mật khẩu", "Người dùng đã đăng nhập", "Người dùng cung cấp mật khẩu cũ và nhập mật khẩu mới hợp lệ để cập nhật.", "Không"],
        ["Quên / Lấy lại MK", "Khách hàng, Organizer", "Nhập email đã đăng ký để nhận mã OTP hoặc đường dẫn thiết lập lại mật khẩu mới.", "<<include>> Gửi link reset mật khẩu qua email"],
        ["Chỉnh sửa thông tin", "Người dùng đã đăng nhập", "Cập nhật thông tin cá nhân: họ tên, ảnh đại diện, số điện thoại liên hệ.", "Không"],
        ["Cài đặt thông báo", "Khách hàng", "Bật/tắt các loại thông báo sự kiện, ưu đãi mở bán vé qua Email hoặc Web Push.", "Không"]
    ]
    add_styled_table(doc, uc1_headers, uc1_data, col_widths=[1.5, 1.2, 2.3, 1.5])

    # Module 2
    add_heading_2(doc, "2.3. Biểu đồ Use Case Phân hệ 2: Tìm kiếm & Xem sự kiện")
    add_body_p(doc, "Phân hệ giúp người dùng khám phá và tìm kiếm sự kiện một cách nhanh chóng, trực quan. Người dùng có thể tìm kiếm theo từ khóa, lọc đa tiêu chí (ngày diễn, địa điểm, thể loại nhạc/kịch/thể thao, mức giá), xem chi tiết thông tin sự kiện cùng dàn nghệ sĩ tham gia, và lưu vào danh sách yêu thích.")
    add_diagram_image(doc, "uc_module2.png", "Hình 2.3: Biểu đồ Use Case Phân hệ 2 - Tìm kiếm & Xem sự kiện")

    uc2_headers = ["Use Case", "Tác nhân", "Mô tả nghiệp vụ", "Quan hệ Use Case"]
    uc2_data = [
        ["Xem sự kiện nổi bật", "Khách hàng", "Hiển thị banner các sự kiện 'Hot', sự kiện sắp mở bán vé hoặc đang bán chạy.", "Không"],
        ["Tìm kiếm sự kiện", "Khách hàng", "Tìm kiếm sự kiện theo tên sự kiện, tên ca sĩ, ban nhạc hoặc địa điểm biểu diễn.", "Không"],
        ["Lọc sự kiện", "Khách hàng", "Lọc theo nhiều tiêu chí kết hợp: Khoảng thời gian, Khu vực (Hà Nội, TP.HCM...), Thể loại, Mức giá.", "Không"],
        ["Xem chi tiết sự kiện", "Khách hàng", "Xem mô tả chi tiết, lịch diễn, danh sách nghệ sĩ, sơ đồ sân khấu và bảng giá vé.", "Không"],
        ["Lưu sự kiện yêu thích", "Khách hàng (đăng nhập)", "Thêm sự kiện vào Wishlist để theo dõi trạng thái và nhận thông báo khi mở bán.", "<<extend>> Nhận thông báo khi mở bán vé"],
        ["Chia sẻ sự kiện", "Khách hàng", "Tạo liên kết chia sẻ sự kiện lên các nền tảng mạng xã hội (Facebook, Zalo, X).", "Không"]
    ]
    add_styled_table(doc, uc2_headers, uc2_data, col_widths=[1.5, 1.2, 2.3, 1.5])

    # Module 3
    add_heading_2(doc, "2.4. Biểu đồ Use Case Phân hệ 3: Xếp hàng và chọn chỗ")
    add_body_p(doc, "Phân hệ trọng yếu đảm bảo hệ thống không bị sập (crash) khi có hàng chục nghìn lượt truy cập đồng thời vào giờ mở bán vé. Cơ chế phòng chờ ảo (Virtual Waiting Room) cấp số thứ tự công bằng, khi tới lượt người mua được chuyển vào màn hình chọn ghế trực quan với đồng hồ đếm ngược giữ ghế tạm thời (10-15 phút).")
    add_diagram_image(doc, "uc_module3.png", "Hình 2.4: Biểu đồ Use Case Phân hệ 3 - Xếp hàng và chọn chỗ")

    uc3_headers = ["Use Case", "Tác nhân", "Mô tả nghiệp vụ", "Quan hệ Use Case"]
    uc3_data = [
        ["Vào hàng đợi ảo", "Khách hàng, Queue System", "Khi sự kiện có tải vượt ngưỡng, hệ thống đưa người mua vào hàng đợi với số thứ tự duy nhất.", "Không"],
        ["Theo dõi tiến trình chờ", "Khách hàng", "Hiển thị số người phía trước và thời gian ước tính; tự động cấp Access Token khi tới lượt.", "Không"],
        ["Xem sơ đồ sân khấu", "Khách hàng", "Tải bản đồ ghế trực quan với màu sắc thể hiện trạng thái: Trống, Đang giữ, Đã bán theo thời gian thực.", "Không"],
        ["Chọn vị trí ghế ngồi", "Khách hàng", "Người dùng nhấn chọn ghế trên sơ đồ (theo khu vực: VIP, Standard, Fan Zone...).", "Không"],
        ["Giữ ghế tạm thời", "Khách hàng, Timer Service", "Khóa các ghế đã chọn trong 10-15 phút, hiển thị đồng hồ đếm ngược để người dùng thanh toán.", "<<extend>> Giải phóng ghế khi hết hạn giữ chỗ"],
        ["Hủy chọn / Thoát chờ", "Khách hàng", "Người dùng chủ động hủy các ghế đang giữ hoặc thoát khỏi hàng đợi.", "Không"]
    ]
    add_styled_table(doc, uc3_headers, uc3_data, col_widths=[1.5, 1.2, 2.3, 1.5])

    # Module 4
    add_heading_2(doc, "2.5. Biểu đồ Use Case Phân hệ 4: Tính tiền và Xuất vé")
    add_body_p(doc, "Phân hệ xử lý thanh toán tài chính và cấp phát vé. Sau khi chọn ghế, đơn hàng được khởi tạo, cho phép nhập mã khuyến mãi, chọn cổng thanh toán phù hợp, hoàn tất thanh toán an toàn và sinh mã vé điện tử QR Code động gửi thẳng tới Email khách hàng.")
    add_diagram_image(doc, "uc_module4.png", "Hình 2.5: Biểu đồ Use Case Phân hệ 4 - Tính tiền và Xuất vé")

    uc4_headers = ["Use Case", "Tác nhân", "Mô tả nghiệp vụ", "Quan hệ Use Case"]
    uc4_data = [
        ["Tạo đơn hàng mua vé", "Khách hàng", "Tổng hợp thông tin các ghế đang giữ, tính tiền tạm tính và tạo mã đơn hàng.", "Không"],
        ["Nhập mã giảm giá", "Khách hàng", "Nhập Voucher ưu đãi; hệ thống kiểm tra điều kiện áp dụng và trừ tiền tương ứng.", "<<include>> Kiểm tra hợp lệ mã voucher"],
        ["Chọn kênh thanh toán", "Khách hàng", "Lựa chọn phương thức: Cổng VNPAY-QR, Ví điện tử MoMo, Thẻ ATM/Visa/MasterCard.", "Không"],
        ["Thanh toán trực tuyến", "Khách hàng, Cổng thanh toán", "Chuyển hướng cổng thanh toán, xác thực OTP ngân hàng và xử lý giao dịch thành công.", "<<include>> Xác thực OTP & Kết quả giao dịch"],
        ["Xuất vé điện tử QR", "Hệ thống, Mail Service", "Sinh mã QR Code được mã hóa chữ ký số (chống làm giả) cho từng ghế trong đơn hàng.", "Không"],
        ["Lưu vào 'Vé của tôi'", "Khách hàng", "Vé điện tử được lưu trữ vĩnh viễn trong tài khoản người dùng để phục vụ soát vé (Check-in).", "Không"]
    ]
    add_styled_table(doc, uc4_headers, uc4_data, col_widths=[1.5, 1.2, 2.3, 1.5])

    # Module 5
    add_heading_2(doc, "2.6. Biểu đồ Use Case Phân hệ 5: Quản lý và thống kê dành cho ban tổ chức")
    add_body_p(doc, "Phân hệ cung cấp bộ công cụ quản trị dành riêng cho Ban tổ chức sự kiện: thiết kế layout khán đài sân khấu, khởi tạo sự kiện mới, định giá vé cho từng khu vực, thiết lập đợt mở bán, quản lý danh sách người tham gia và xem báo cáo tổng kết doanh thu trực quan.")
    add_diagram_image(doc, "uc_module5.png", "Hình 2.6: Biểu đồ Use Case Phân hệ 5 - Quản lý & Thống kê dành cho BTC")

    uc5_headers = ["Use Case", "Tác nhân", "Mô tả nghiệp vụ", "Quan hệ Use Case"]
    uc5_data = [
        ["Thiết kế sơ đồ ghế", "Ban tổ chức", "Công cụ vẽ khán đài trực quan: thiết lập khu vực ghế ngồi, hàng ghế, số ghế, vé đứng tự do.", "Không"],
        ["Tạo sự kiện mới", "Ban tổ chức, Admin", "Nhập thông tin sự kiện: tiêu đề, poster, thể loại, thông tin nghệ sĩ, địa điểm tổ chức.", "Không"],
        ["Cấu hình giá vé & Suất diễn", "Ban tổ chức", "Thiết lập các khung giờ diễn và gán bảng giá vé riêng cho từng khu vực ghế (VIP, Regular).", "Không"],
        ["Quản lý mở bán & Voucher", "Ban tổ chức", "Tạo các đợt mở bán (Early Bird, General Sale) và phát hành mã giảm giá kèm ngân sách.", "Không"],
        ["Quản lý check-in soát vé", "Ban tổ chức", "Ứng dụng quét mã QR tại cổng ra vào để xác thực vé hợp lệ và ghi nhận thời gian vào cửa.", "Không"],
        ["Báo cáo thống kê doanh thu", "Ban tổ chức, Admin", "Xem tổng quan doanh thu, số vé đã bán, tỷ lệ lấp đầy khán đài theo thời gian thực.", "<<extend>> Xuất báo cáo (Excel/PDF)"]
    ]
    add_styled_table(doc, uc5_headers, uc5_data, col_widths=[1.5, 1.2, 2.3, 1.5])

    # =========================================================================
    # PHẦN 3
    # =========================================================================
    add_heading_1(doc, "3. BIỂU ĐỒ LỚP THỰC THỂ PHÂN TÍCH CHO HỆ THỐNG VÀ CHO CÁC CHỨC NĂNG")
    
    add_body_p(doc, "Biểu đồ lớp thực thể phân tích (Analysis Entity Class Diagram) là mô hình hướng khái niệm (Conceptual Model), tập trung thể hiện các thực thể thông tin nghiệp vụ cốt lõi của miền bài toán. Theo chuẩn phân tích:", "Nguyên lý xây dựng: ")
    add_bullet_p(doc, "Chỉ bao gồm Tên thực thể và Danh sách thuộc tính nghiệp vụ thuần túy.", "Cấu trúc lớp phân tích: ")
    add_bullet_p(doc, "Không chứa các chi tiết kỹ thuật lập trình như kiểu dữ liệu cụ thể (int, varchar2...), không chứa các phương thức get/set kỹ thuật.", "Độc lập công nghệ: ")
    add_bullet_p(doc, "Tập trung thể hiện mối quan hệ ngữ nghĩa: Kế thừa (Generalization), Kết hợp (Association), Kết tập/Bao hàm (Composition/Aggregation) kèm bản số (Multiplicity: 1, 1..*, 0..1, *).", "Mối quan hệ: ")

    add_heading_2(doc, "3.1. Biểu đồ lớp thực thể phân tích tổng thể hệ thống")
    add_body_p(doc, "Biểu đồ bao quát toàn bộ các thực thể phân tích nòng cốt của hệ thống (được đồng bộ chuẩn hóa tên tiếng Anh theo mô hình thiết kế) và các mối liên kết nghiệp vụ giữa chúng:")
    add_diagram_image(doc, "analysis_tong_the.png", "Hình 3.1: Biểu đồ lớp thực thể phân tích tổng thể toàn hệ thống")

    ana_headers = ["Tên lớp thực thể (Class)", "Ý nghĩa nghiệp vụ", "Các thuộc tính phân tích (Attributes)"]
    ana_data = [
        ["User", "Thông tin tài khoản chung trong hệ thống", "userId, username, password, fullName, email, phoneNumber, role"],
        ["Customer", "Kế thừa User, đại diện người mua vé", "rewardPoints, memberTier"],
        ["Organizer", "Kế thừa User, đơn vị tổ chức sự kiện", "companyName, taxId, businessLicense, contactInfo"],
        ["Venue", "Nơi diễn ra sự kiện (sân vận động, nhà hát)", "venueId, venueName, address, capacity"],
        ["SeatMap", "Bản thiết kế phân bổ chỗ ngồi của địa điểm", "seatMapId, mapName, stageType, totalSeats"],
        ["SeatZone", "Khu vực khán đài cụ thể (VIP, VVIP, CAT 1)", "zoneId, zoneName, colorHex, zoneType, seatCount"],
        ["Seat", "Vị trí ghế cụ thể trên khán đài", "seatId, rowLabel, seatNumber, coordinateX, coordinateY"],
        ["Event", "Sự kiện được tổ chức (Concert, liveshow, kịch)", "eventId, title, category, description, bannerUrl, status"],
        ["Showtime", "Một buổi biểu diễn cụ thể của sự kiện", "showtimeId, showDate, startTime, endTime"],
        ["ZonePricing", "Định giá vé cho từng khu vực của suất diễn", "pricingId, price, maxQuota, soldCount"],
        ["VirtualQueue", "Bản ghi xếp hàng chờ của khách hàng", "queueId, queueNumber, accessToken, status"],
        ["SeatHold", "Trạng thái khóa giữ ghế trong 10-15 phút", "holdId, holdToken, startTime, expiresAt"],
        ["Order", "Đơn hàng thanh toán vé của khách hàng", "orderId, orderCode, subTotal, discount, finalTotal, status"],
        ["Ticket", "Vé điện tử cấp phát cho từng chỗ ngồi", "ticketId, ticketCode, qrCodeHash, price, isUsed"],
        ["Voucher", "Mã giảm giá khuyến mãi áp dụng cho đơn hàng", "voucherId, code, percentOff, maxDiscount, validUntil"],
        ["PaymentTransaction", "Giao dịch thanh toán trực tuyến qua cổng", "transactionId, transRef, paymentMethod, transAmount, status"]
    ]
    add_styled_table(doc, ana_headers, ana_data, col_widths=[1.5, 2.3, 2.7])

    # Module 1 Analysis
    add_heading_2(doc, "3.2. Biểu đồ lớp thực thể phân tích Phân hệ 1: Quản lý tài khoản")
    add_body_p(doc, "Thể hiện quan hệ kế thừa giữa User với Customer và Organizer, đồng thời liên kết với UserSession (phiên làm việc) và Notification (thông báo gửi tới người dùng).")
    add_diagram_image(doc, "analysis_module1.png", "Hình 3.2: Biểu đồ lớp thực thể phân tích - Phân hệ 1: Quản lý tài khoản")

    # Module 2 Analysis
    add_heading_2(doc, "3.3. Biểu đồ lớp thực thể phân tích Phân hệ 2: Tìm kiếm & Xem sự kiện")
    add_body_p(doc, "Mô hình hóa thực thể Event liên kết với Category, Venue tổ chức, dàn Artist biểu diễn và danh sách FavoriteEvent của khách hàng.")
    add_diagram_image(doc, "analysis_module2.png", "Hình 3.3: Biểu đồ lớp thực thể phân tích - Phân hệ 2: Tìm kiếm & Xem sự kiện")

    # Module 3 Analysis
    add_heading_2(doc, "3.4. Biểu đồ lớp thực thể phân tích Phân hệ 3: Xếp hàng và chọn chỗ")
    add_body_p(doc, "Mô hình hóa cấu trúc phân cấp không gian từ SeatMap -> SeatZone -> Seat, kết hợp với VirtualQueue và thực thể SeatHold khóa giữ các ghế được chọn.")
    add_diagram_image(doc, "analysis_module3.png", "Hình 3.4: Biểu đồ lớp thực thể phân tích - Phân hệ 3: Xếp hàng và chọn chỗ")

    # Module 4 Analysis
    add_heading_2(doc, "3.5. Biểu đồ lớp thực thể phân tích Phân hệ 4: Tính tiền và Xuất vé")
    add_body_p(doc, "Order đóng vai trò trung tâm liên kết với Voucher (0..1), PaymentTransaction (1) và chứa tập hợp các Ticket (1..*), mỗi Ticket tương ứng một Seat duy nhất.")
    add_diagram_image(doc, "analysis_module4.png", "Hình 3.5: Biểu đồ lớp thực thể phân tích - Phân hệ 4: Tính tiền và Xuất vé")

    # Module 5 Analysis
    add_heading_2(doc, "3.6. Biểu đồ lớp thực thể phân tích Phân hệ 5: Quản lý và thống kê BTC")
    add_body_p(doc, "Organizer quản lý nhiều Event, mỗi sự kiện áp dụng SeatMap và nhiều Showtime, từ đó thiết lập ZonePricing và kết xuất RevenueReport.")
    add_diagram_image(doc, "analysis_module5.png", "Hình 3.6: Biểu đồ lớp thực thể phân tích - Phân hệ 5: Quản lý & Thống kê BTC")

    add_heading_3(doc, "3.6.1. Biểu đồ lớp phân tích chi tiết cho chức năng Xem thống kê doanh thu sự kiện")
    add_body_p(doc, "Để đáp ứng nguyên tắc hướng đối tượng cốt lõi (Lớp phải đóng gói trạng thái và hành vi tính toán nghiệp vụ, không chỉ là tập hợp dữ liệu thụ động/Anemic Domain Model), biểu đồ lớp phân tích dưới đây thể hiện chi tiết các thực thể tham gia ca sử dụng đơn lẻ 'Xem thống kê doanh thu sự kiện' cùng các phương thức tính toán tài chính trọng tâm:")
    add_diagram_image(doc, "analysis_m5_view_revenue_detail.png", "Hình 3.7: Biểu đồ lớp phân tích chi tiết - Chức năng Xem thống kê doanh thu sự kiện")

    # =========================================================================
    # PHẦN 4
    # =========================================================================
    add_heading_1(doc, "4. BIỂU ĐỒ LỚP THỰC THỂ THIẾT KẾ CHO HỆ THỐNG VÀ CHO CÁC CHỨC NĂNG")
    
    add_body_p(doc, "Biểu đồ lớp thực thể thiết kế (Design Entity Class Diagram) là mô hình mức thiết kế chi tiết (Detailed Object-Oriented Design), sẵn sàng cho việc sinh mã nguồn trong các ngôn ngữ lập trình hiện đại (Java Spring Boot, C# .NET Core, TypeScript) và ánh xạ quan hệ ORM (Hibernate / JPA).", "Nguyên lý xây dựng: ")
    add_bullet_p(doc, "Stereotype <<entity>> định danh các lớp thực thể lưu trữ dữ liệu.", "Stereotype chuẩn: ")
    add_bullet_p(doc, "Đầy đủ 3 ngăn chuẩn UML: Tên lớp, Thuộc tính có phạm vi truy cập (- private) và kiểu dữ liệu cụ thể (Long, String, LocalDateTime, double, boolean...), Phương thức có phạm vi (+ public) và danh sách tham số.", "Cấu trúc 3 ngăn chi tiết: ")
    add_bullet_p(doc, "Xác định rõ các thuộc tính định danh khóa chính (Primary Key - id: Long) và các ràng buộc toàn vẹn.", "Khóa chính và ràng buộc: ")
    add_bullet_p(doc, "Bao gồm các phương thức nghiệp vụ nội tại của đối tượng (ví dụ: isAvailable(), lock(), calculateTotal(), generateQR()).", "Phương thức nghiệp vụ: ")

    add_heading_2(doc, "4.1. Biểu đồ lớp thực thể thiết kế tổng thể hệ thống")
    add_body_p(doc, "Biểu đồ lớp thực thể thiết kế tổng thể thể hiện mối liên kết chặt chẽ giữa tất cả các thực thể nghiệp vụ trong hệ thống:")
    add_diagram_image(doc, "design_tong_the.png", "Hình 4.1: Biểu đồ lớp thực thể thiết kế tổng thể toàn hệ thống")

    # Module 1 Design
    add_heading_2(doc, "4.2. Biểu đồ lớp thực thể thiết kế Phân hệ 1: Quản lý tài khoản")
    add_body_p(doc, "Mô hình thiết kế chi tiết các lớp User, Customer, Organizer, UserSession và Notification:")
    add_diagram_image(doc, "design_module1.png", "Hình 4.2: Biểu đồ lớp thực thể thiết kế - Phân hệ 1: Quản lý tài khoản")

    d1_headers = ["Lớp thiết kế", "Thuộc tính chi tiết (Kiểu dữ liệu & Ràng buộc)", "Phương thức nghiệp vụ (Operations)"]
    d1_data = [
        ["User\n(Base Entity)", "- id: Long [PK, Not Null]\n- username: String [Unique]\n- passwordHash: String [BCrypt]\n- email: String [Unique]\n- phone: String\n- role: UserRole [Enum: CUSTOMER, ORGANIZER, ADMIN]", "+ authenticate(pwd: String): boolean\n+ changePassword(newPwd: String): void\n+ updateInfo(dto: UserDto): void\n+ isAccountNonLocked(): boolean"],
        ["Customer\n(extends User)", "- rewardPoints: int [Default: 0]\n- memberTier: MemberTier [BRONZE, SILVER, GOLD, VIP]", "+ addRewardPoints(pts: int): void\n+ getOrderHistory(): List<Order>\n+ canRedeemPoints(pts: int): boolean"],
        ["Organizer\n(extends User)", "- orgName: String [Not Null]\n- taxId: String [Unique]\n- businessLicense: String", "+ verifyLicense(): boolean\n+ getManagedEvents(): List<Event>\n+ updateOrgProfile(info: OrgDto): void"],
        ["UserSession", "- id: Long [PK]\n- jwtToken: String [Not Null]\n- createdAt: LocalDateTime\n- expiresAt: LocalDateTime", "+ isExpired(): boolean\n+ refreshToken(): String\n+ revoke(): void"],
        ["Notification", "- id: Long [PK]\n- title: String\n- message: String\n- isRead: boolean [Default: false]\n- sentAt: LocalDateTime", "+ markAsRead(): void\n+ sendEmailNotification(): boolean"]
    ]
    add_styled_table(doc, d1_headers, d1_data, col_widths=[1.5, 2.5, 2.5])

    # Module 2 Design
    add_heading_2(doc, "4.3. Biểu đồ lớp thực thể thiết kế Phân hệ 2: Tìm kiếm & Xem sự kiện")
    add_body_p(doc, "Mô hình thiết kế chi tiết các lớp Event, Showtime, Category, Venue, Artist và FavoriteEvent:")
    add_diagram_image(doc, "design_module2.png", "Hình 4.3: Biểu đồ lớp thực thể thiết kế - Phân hệ 2: Tìm kiếm & Xem sự kiện")

    d2_headers = ["Lớp thiết kế", "Thuộc tính chi tiết (Kiểu dữ liệu & Ràng buộc)", "Phương thức nghiệp vụ (Operations)"]
    d2_data = [
        ["Event", "- id: Long [PK]\n- title: String [Not Null]\n- description: String [Text]\n- bannerUrl: String\n- status: EventStatus [DRAFT, PUBLISHED, SOLD_OUT, ENDED]", "+ publish(): void\n+ cancel(): void\n+ isUpcoming(): boolean\n+ getActiveShowtimes(): List<Showtime>"],
        ["Showtime", "- id: Long [PK]\n- showDate: LocalDate [Not Null]\n- startTime: LocalTime [Not Null]\n- endTime: LocalTime", "+ isAvailableForBooking(): boolean\n+ getAvailableSeats(): int\n+ isOngoing(): boolean"],
        ["Category", "- id: Long [PK]\n- name: String [Unique, Not Null]\n- code: String [Unique]", "+ getEvents(): List<Event>\n+ countEvents(): int"],
        ["Venue", "- id: Long [PK]\n- name: String [Not Null]\n- address: String [Not Null]\n- capacity: int [> 0]", "+ checkCapacity(): boolean\n+ getAssignedSeatMap(): SeatMap"],
        ["Artist", "- id: Long [PK]\n- stageName: String [Not Null]\n- bio: String\n- avatarUrl: String", "+ getParticipatingEvents(): List<Event>"],
        ["FavoriteEvent", "- id: Long [PK]\n- customerId: Long [FK]\n- savedDate: LocalDateTime", "+ notifySaleOpening(): void\n+ remove(): void"]
    ]
    add_styled_table(doc, d2_headers, d2_data, col_widths=[1.5, 2.5, 2.5])

    # Module 3 Design
    add_heading_2(doc, "4.4. Biểu đồ lớp thực thể thiết kế Phân hệ 3: Xếp hàng và chọn chỗ")
    add_body_p(doc, "Mô hình thiết kế chi tiết các lớp SeatMap, SeatZone, Seat, VirtualQueue và SeatHold:")
    add_diagram_image(doc, "design_module3.png", "Hình 4.4: Biểu đồ lớp thực thể thiết kế - Phân hệ 3: Xếp hàng và chọn chỗ")

    d3_headers = ["Lớp thiết kế", "Thuộc tính chi tiết (Kiểu dữ liệu & Ràng buộc)", "Phương thức nghiệp vụ (Operations)"]
    d3_data = [
        ["SeatMap", "- id: Long [PK]\n- mapName: String [Not Null]\n- totalSeats: int [> 0]\n- layoutSvg: String [SVG Schema]", "+ validateLayout(): boolean\n+ getZones(): List<SeatZone>\n+ getTotalCapacity(): int"],
        ["SeatZone", "- id: Long [PK]\n- zoneName: String [VIP, CAT 1...]\n- colorHex: String [#FF0000]\n- zoneType: ZoneType [SEATED, STANDING]", "+ getSeats(): List<Seat>\n+ countVacantSeats(): int\n+ isStandingZone(): boolean"],
        ["Seat", "- id: Long [PK]\n- rowLabel: String [A, B, C...]\n- seatNumber: int [1, 2, 3...]\n- isLocked: boolean [Default: false]", "+ lockSeat(token: String): boolean\n+ unlockSeat(): void\n+ isAvailable(): boolean\n+ getSeatCode(): String"],
        ["VirtualQueue", "- id: Long [PK]\n- queueNumber: int [Auto-increment]\n- accessToken: String [UUID]\n- status: QueueStatus [WAITING, SERVING, EXPIRED]", "+ isTurn(): boolean\n+ generateToken(): String\n+ expireToken(): void"],
        ["SeatHold", "- id: Long [PK]\n- holdToken: String [UUID, Unique]\n- startTime: LocalDateTime\n- expiresAt: LocalDateTime", "+ isExpired(): boolean\n+ getRemainingSeconds(): long\n+ confirmHold(): void\n+ releaseHold(): void"]
    ]
    add_styled_table(doc, d3_headers, d3_data, col_widths=[1.5, 2.5, 2.5])

    # Module 4 Design
    add_heading_2(doc, "4.5. Biểu đồ lớp thực thể thiết kế Phân hệ 4: Tính tiền và Xuất vé")
    add_body_p(doc, "Mô hình thiết kế chi tiết các lớp Order, Voucher, PaymentTransaction và Ticket:")
    add_diagram_image(doc, "design_module4.png", "Hình 4.5: Biểu đồ lớp thực thể thiết kế - Phân hệ 4: Tính tiền và Xuất vé")

    d4_headers = ["Lớp thiết kế", "Thuộc tính chi tiết (Kiểu dữ liệu & Ràng buộc)", "Phương thức nghiệp vụ (Operations)"]
    d4_data = [
        ["Order", "- id: Long [PK]\n- orderCode: String [ORD-XXXXX]\n- subTotal: double [>= 0]\n- discount: double [>= 0]\n- finalTotal: double [>= 0]\n- status: OrderStatus [PENDING, PAID, CANCELLED]", "+ calculateFinalTotal(): double\n+ applyVoucher(v: Voucher): boolean\n+ markAsPaid(): void\n+ cancelOrder(): void"],
        ["Voucher", "- id: Long [PK]\n- code: String [Unique, Not Null]\n- percentOff: double [0.0 - 100.0]\n- maxDiscount: double\n- validUntil: LocalDate", "+ isValid(amount: double): boolean\n+ calculateDiscount(amount: double): double\n+ decrementQuota(): void"],
        ["PaymentTransaction", "- id: Long [PK]\n- transRef: String [Bank Reference]\n- method: PaymentMethod [VNPAY, MOMO, VISA]\n- transAmount: double\n- status: TransStatus [SUCCESS, FAILED, PENDING]", "+ processPayment(): boolean\n+ verifyWebhook(payload: String): boolean\n+ getGatewayUrl(): String"],
        ["Ticket", "- id: Long [PK]\n- ticketCode: String [TCK-XXXXX]\n- qrCodeHash: String [HMAC-SHA256 Encrypted]\n- price: double\n- isUsed: boolean [Default: false]\n- checkinTime: LocalDateTime", "+ generateEncryptedQR(): String\n+ checkIn(): boolean\n+ isCheckedIn(): boolean\n+ cancelTicket(): void"]
    ]
    add_styled_table(doc, d4_headers, d4_data, col_widths=[1.5, 2.5, 2.5])

    # Module 5 Design
    add_heading_2(doc, "4.6. Biểu đồ lớp thực thể thiết kế Phân hệ 5: Quản lý và thống kê BTC")
    add_body_p(doc, "Mô hình thiết kế chi tiết các lớp Organizer, Event, SeatMap, Showtime, ZonePricing và RevenueReport:")
    add_diagram_image(doc, "design_module5.png", "Hình 4.6: Biểu đồ lớp thực thể thiết kế - Phân hệ 5: Quản lý & Thống kê BTC")

    d5_headers = ["Lớp thiết kế", "Thuộc tính chi tiết (Kiểu dữ liệu & Ràng buộc)", "Phương thức nghiệp vụ (Operations)"]
    d5_data = [
        ["Organizer", "- id: Long [PK]\n- companyName: String\n- taxId: String [Unique]\n- email: String", "+ createEvent(dto: EventDto): Event\n+ updatePricing(pricingDto): void\n+ viewReports(eventId: Long): RevenueReport"],
        ["ZonePricing", "- id: Long [PK]\n- price: double [>= 0]\n- maxQuota: int [Total Allocated Seats]\n- soldCount: int [Current Sold Seats]", "+ calculateZoneGross(): double\n+ calculateRemainingSeats(): int\n+ getSoldRate(): double\n+ isSoldOut(): boolean\n+ recordSale(quantity: int): void"],
        ["RevenueReport", "- id: Long [PK]\n- eventId: Long\n- calculatedAt: LocalDateTime\n- taxRate: double\n- platformCommissionRate: double", "+ calculateTotalCapacity(): int\n+ calculateTotalTicketsSold(): int\n+ calculateGrossRevenue(): double\n+ calculateNetRevenue(): double\n+ calculateOccupancyRate(): double\n+ calculateAverageTicketPrice(): double\n+ calculateZoneContribution(zoneId: Long): double\n+ evaluatePerformance(): PerformanceStatus\n+ exportExcel(): byte[]\n+ exportPdf(): byte[]"]
    ]
    add_styled_table(doc, d5_headers, d5_data, col_widths=[1.5, 2.5, 2.5])

    # Module 5 Key Use Case Detailed Design
    add_heading_2(doc, "4.7. Thiết kế chi tiết chức năng trọng tâm Module 5: Xem thống kê doanh thu sự kiện")
    add_body_p(doc, "Trong Module 5, chức năng 'Xem thống kê doanh thu sự kiện' (View Event Revenue Statistics) là một chức năng nghiệp vụ đơn lẻ (Single Atomic Function - Read/Calculate Analytics) đóng vai trò quyết định hiệu quả kinh doanh của Ban tổ chức. Chức năng này không đơn thuần là truy vấn dữ liệu thô (DTO) mà đòi hỏi các lớp thực thể phải sở hữu các phương thức thực hiện tính toán tài chính phức tạp, đảm bảo tính đóng gói (Encapsulation) chuẩn mực của lập trình hướng đối tượng.", "Lý do lựa chọn chức năng: ")

    add_heading_3(doc, "a. Đặc tả ca sử dụng chi tiết (Use Case Specification)")
    uc_spec_headers = ["Thuộc tính đặc tả", "Nội dung chi tiết"]
    uc_spec_data = [
        ["Tên chức năng / Use Case", "Xem thống kê doanh thu sự kiện (View Event Revenue Statistics)"],
        ["Phân hệ trực thuộc", "Module 5: Quản lý và thống kê dành cho Ban tổ chức"],
        ["Tác nhân (Actor)", "Ban tổ chức sự kiện (Organizer), Quản trị viên (Admin)"],
        ["Mục tiêu nghiệp vụ", "Cung cấp bức tranh tài chính toàn cảnh của sự kiện: doanh thu gộp, doanh thu thuần sau thuế và phí sàn, tỷ lệ lấp đầy khán đài, giá vé bình quân và xếp hạng hiệu suất mở bán."],
        ["Tiền điều kiện (Pre-conditions)", "1. Ban tổ chức đã đăng nhập thành công vào hệ thống.\n2. Sự kiện đã được cấu hình sơ đồ ghế và đã phát sinh giao dịch bán vé."],
        ["Hậu điều kiện (Post-conditions)", "Hệ thống tổng hợp và hiển thị trực quan các chỉ số tài chính (KPIs) cùng biểu đồ phân bổ doanh thu theo từng khu vực khán đài."],
        ["Luồng sự kiện chính (Main Flow)", "1. Ban tổ chức chọn sự kiện cần xem từ danh sách sự kiện do mình quản lý.\n2. Giao diện RevenueReportView gửi yêu cầu tra cứu tới RevenueReportController.\n3. Controller khởi tạo đối tượng RevenueReport nạp dữ liệu cấu hình vé và các giao dịch.\n4. RevenueReport thực hiện chuỗi phương thức tính toán nội tại: calculateGrossRevenue(), calculateNetRevenue(), calculateOccupancyRate(), evaluatePerformance().\n5. Controller đóng gói kết quả vào RevenueSummaryDto và trả về cho RevenueReportView.\n6. Giao diện kết xuất các thẻ KPI và biểu đồ phân bổ doanh thu trực quan."]
    ]
    add_styled_table(doc, uc_spec_headers, uc_spec_data, col_widths=[2.2, 4.3])

    add_heading_3(doc, "b. Biểu đồ tuần tự (Sequence Diagram) thể hiện tương tác BCE")
    add_body_p(doc, "Biểu đồ tuần tự thể hiện sự tương tác mạch lạc giữa tác nhân Ban tổ chức, lớp Boundary (Giao diện), lớp Control (Điều phối) và các lớp Entity (Thực thể tính toán):")
    add_diagram_image(doc, "seq_m5_view_revenue.png", "Hình 4.7: Biểu đồ tuần tự ca sử dụng Xem thống kê doanh thu sự kiện")

    add_heading_3(doc, "c. Biểu đồ lớp thiết kế chi tiết theo mô hình BCE")
    add_body_p(doc, "Mô hình thiết kế 3 lớp (Boundary - Control - Entity) làm nổi bật các phương thức tính toán nghiệp vụ trong các lớp thực thể RevenueReport và ZonePricing:")
    add_diagram_image(doc, "design_m5_view_revenue_detail.png", "Hình 4.8: Biểu đồ lớp thiết kế chi tiết (BCE) chức năng Thống kê doanh thu")

    add_heading_3(doc, "d. Bảng phân tích chi tiết các phương thức tính toán (Computational Methods)")
    add_body_p(doc, "Để các lớp không bị biến thành 'cấu trúc dữ liệu thụ động' (Anemic Domain Model), toàn bộ logic tính toán tài chính được đóng gói trực tiếp vào các thực thể:")

    calc_headers = ["Tên phương thức tính toán", "Lớp sở hữu", "Kiểu trả về", "Công thức / Thuật toán tính toán"]
    calc_data = [
        ["calculateTotalCapacity()", "RevenueReport", "int", "Tổng sức chứa của khán đài = tổng số ghế của tất cả các phân khu (ZonePricing):\nTotalCapacity = sum(zone.seatCount)"],
        ["calculateTotalTicketsSold()", "RevenueReport", "int", "Tổng số vé thực tế đã bán thành công = tổng số vé bán ra của từng phân khu:\nTotalSold = sum(zone.soldCount)"],
        ["calculateGrossRevenue()", "RevenueReport", "double", "Tổng doanh thu gộp = tổng tích số giữa số vé đã bán và đơn giá vé từng khu vực:\nGrossRevenue = sum(zone.soldCount * zone.price)"],
        ["calculateNetRevenue()", "RevenueReport", "double", "Doanh thu thuần thực nhận sau khi khấu trừ thuế VAT và phí nền tảng:\nNetRevenue = GrossRevenue * (1 - taxRate - platformFeeRate)"],
        ["calculateOccupancyRate()", "RevenueReport", "double", "Tỷ lệ lấp đầy khán đài theo phần trăm:\nOccupancyRate = (TotalSold / (double) TotalCapacity) * 100.0%"],
        ["calculateAverageTicketPrice()", "RevenueReport", "double", "Giá vé bình quân trên mỗi vé bán ra:\nAvgPrice = TotalSold > 0 ? (GrossRevenue / TotalSold) : 0.0"],
        ["calculateZoneContribution(zoneId)", "RevenueReport", "double", "Tỷ trọng đóng góp doanh thu của một khu vực so với tổng doanh thu:\nContribution = (ZoneGross / GrossRevenue) * 100.0%"],
        ["evaluatePerformance()", "RevenueReport", "PerformanceStatus", "Đánh giá phân loại sức mua theo thang tỷ lệ lấp đầy:\n- Occupancy >= 85%: EXCELLENT (Cháy vé)\n- 70% <= Occupancy < 85%: GOOD (Đạt chỉ tiêu)\n- 50% <= Occupancy < 70%: AVERAGE (Hòa vốn)\n- Occupancy < 50%: POOR (Cần giải cứu vé)"],
        ["calculateZoneGross()", "ZonePricing", "double", "Doanh thu của riêng phân khu vé = price * soldCount"],
        ["getSoldRate()", "ZonePricing", "double", "Tỷ lệ bán của riêng phân khu = (soldCount / (double) maxQuota) * 100.0%"]
    ]
    add_styled_table(doc, calc_headers, calc_data, col_widths=[1.8, 1.1, 1.1, 2.5])

    # Final summary conclusion
    add_heading_1(doc, "KẾT LUẬN VÀ CAM KẾT HOÀN THÀNH")
    add_body_p(doc, "Tài liệu phân tích và thiết kế hệ thống trên đã hoàn thiện đầy đủ 4 yêu cầu kiểm tra môn học theo đúng đề cương bài tập lớn do Thầy Nguyễn Đức Hiển giao:")
    add_bullet_p(doc, "Đầy đủ thông tin nhóm, đề tài, phân công nhiệm vụ và phạm vi 5 module chức năng cho các thành viên.", "1. Thông tin nhóm: ")
    add_bullet_p(doc, "Biểu đồ Use Case tổng thể toàn hệ thống và 5 biểu đồ Use Case phân rã chi tiết cho từng phân hệ kèm bảng đặc tả tóm tắt nghiệp vụ.", "2. Biểu đồ Use Case: ")
    add_bullet_p(doc, "Biểu đồ lớp thực thể phân tích tổng thể và 5 biểu đồ lớp thực thể phân tích phân rã theo 5 module, thể hiện đúng bản chất mô hình hóa khái niệm miền bài toán.", "3. Biểu đồ lớp phân tích: ")
    add_bullet_p(doc, "Biểu đồ lớp thực thể thiết kế tổng thể và 5 biểu đồ lớp thực thể thiết kế chi tiết theo 5 module với cấu trúc 3 ngăn chuẩn UML, kiểu dữ liệu chặt chẽ, ràng buộc khóa chính và phương thức nghiệp vụ.", "4. Biểu đồ lớp thiết kế: ")

    doc.save(DOCX_OUT)
    print(f"Document successfully generated and saved to: {DOCX_OUT}")

if __name__ == "__main__":
    build_full_docx()
